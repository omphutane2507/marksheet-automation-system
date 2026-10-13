import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from marks.models import ClassRoom, Student, Subject, Mark, GradeRule
from marks.services.calculations import (
    calculate_student_consolidated,
    calculate_subject_analysis,
    calculate_class_overview
)


def export_class_excel(classroom):
    """
    Generates a formatted multi-sheet Excel workbook for a ClassRoom matching the sample layout:
    - One sheet per subject (Subject data table + Subject Result Analysis table)
    - Consolidated sheet (Student marks in all subjects, totals, total %, overall grade, class grade count)
    - Overall Result Analysis sheet (Side-by-side grade counts comparison per subject)
    
    Returns: BytesIO buffer containing the Excel binary content.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # remove default sheet

    # Fetch database data
    students = list(Student.objects.filter(classroom=classroom).order_by('roll_number'))
    subjects = list(Subject.objects.filter(classroom=classroom).order_by('order', 'name'))
    grade_rules = list(GradeRule.objects.all())

    # Map student subject marks: { student_id: { subject_id: mark_obj } }
    marks_query = Mark.objects.filter(student__classroom=classroom).select_related('student', 'subject')
    marks_map = {}
    for m in marks_query:
        if m.student_id not in marks_map:
            marks_map[m.student_id] = {}
        marks_map[m.student_id][m.subject_id] = m

    # Styles
    font_title = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    font_bold = Font(name='Calibri', size=10, bold=True)
    font_regular = Font(name='Calibri', size=10)

    fill_header = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid') # Navy blue
    fill_sub_header = PatternFill(start_color='3B82F6', end_color='3B82F6', fill_type='solid') # Accent blue
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # -------------------------------------------------------------
    # 1. CREATE SUBJECT SHEETS
    # -------------------------------------------------------------
    subject_analysis_results = {}

    for subject in subjects:
        ws = wb.create_sheet(title=subject.name[:30]) # max title length 31
        ws.views.sheetView[0].showGridLines = True

        # Title Block
        ws.merge_cells('A1:G1')
        title_cell = ws['A1']
        title_cell.value = f"{classroom.class_name} {classroom.division} ({classroom.academic_year.name}) - {subject.name} Result Sheet"
        title_cell.font = font_title
        title_cell.fill = fill_header
        title_cell.alignment = align_center

        # Header Row
        headers = ['SR.NO.', 'ROLL NO', 'NAME OF STUDENT', 'MARKS OBTAINED', 'OUT OF', 'PERCENTAGE', 'GRADE']
        ws.append([]) # Row 2 blank
        ws.append(headers) # Row 3

        for col_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx)
            cell.font = font_header
            cell.fill = fill_sub_header
            cell.alignment = align_center
            cell.border = thin_border

        # Student Mark Data Rows
        subject_mark_records = []

        for idx, student in enumerate(students, 1):
            mark_obj = marks_map.get(student.id, {}).get(subject.id)
            is_absent = mark_obj.is_absent if mark_obj else False
            marks_val = mark_obj.marks_obtained if mark_obj else None

            subject_mark_records.append({
                'marks_obtained': marks_val,
                'is_absent': is_absent
            })

            # Calculate individual subject result
            if is_absent or marks_val is None:
                m_str = "AB"
                p_str = "N/A"
                g_str = "AB"
            else:
                m_str = marks_val
                p_val = round((marks_val / subject.out_of_marks) * 100.0, 2)
                p_str = f"{p_val:.2f}%"
                from marks.services.calculations import determine_grade
                g_str = determine_grade(p_val, grade_rules)

            row_data = [idx, student.roll_number, student.name, m_str, subject.out_of_marks, p_str, g_str]
            ws.append(row_data)

            current_row = idx + 3
            for c_idx in range(1, 8):
                cell = ws.cell(row=current_row, column=c_idx)
                cell.font = font_regular
                cell.border = thin_border
                if c_idx in [1, 2, 4, 5, 6, 7]:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        # Subject Analysis Table at Bottom
        sub_analysis = calculate_subject_analysis(subject.name, subject.out_of_marks, subject_mark_records, classroom.pass_percentage, grade_rules)
        subject_analysis_results[subject.id] = sub_analysis

        analysis_start_row = len(students) + 6
        ws.cell(row=analysis_start_row, column=1, value="SUBJECT RESULT ANALYSIS").font = font_bold

        a_headers = ['Grade A', 'Grade B', 'Grade C', 'Grade D', 'Grade E', 'Absent (AB)', 'Total Students', 'Subject Avg Marks', 'Pass %']
        for c_idx, ah in enumerate(a_headers, 1):
            cell = ws.cell(row=analysis_start_row + 1, column=c_idx)
            cell.value = ah
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        gc = sub_analysis['grade_counts']
        a_values = [
            gc.get('A', 0), gc.get('B', 0), gc.get('C', 0), gc.get('D', 0), gc.get('E', 0), gc.get('AB', 0),
            sub_analysis['total_students'], sub_analysis['subject_avg_marks'], f"{sub_analysis['subject_pass_percent']}%"
        ]
        for c_idx, av in enumerate(a_values, 1):
            cell = ws.cell(row=analysis_start_row + 2, column=c_idx)
            cell.value = av
            cell.font = font_bold
            cell.alignment = align_center
            cell.border = thin_border

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # -------------------------------------------------------------
    # 2. CONSOLIDATED SHEET
    # -------------------------------------------------------------
    ws_con = wb.create_sheet(title="Consolidated")
    ws_con.views.sheetView[0].showGridLines = True

    # Title Block
    num_cols = 5 + len(subjects) + 4
    ws_con.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    title_cell = ws_con.cell(row=1, column=1)
    title_cell.value = f"{classroom.class_name} {classroom.division} ({classroom.academic_year.name}) - Consolidated Marks Sheet"
    title_cell.font = font_title
    title_cell.fill = fill_header
    title_cell.alignment = align_center

    # Header Row
    c_headers = ['SR.NO.', 'ROLL NO', 'NAME OF STUDENT'] + [s.name.upper() for s in subjects] + ['TOTAL OBTAINED', 'TOTAL OUT OF', 'TOTAL %', 'OVERALL GRADE', 'RESULT']
    ws_con.append([])
    ws_con.append(c_headers)

    for c_idx, ch in enumerate(c_headers, 1):
        cell = ws_con.cell(row=3, column=c_idx)
        cell.font = font_header
        cell.fill = fill_sub_header
        cell.alignment = align_center
        cell.border = thin_border

    class_student_results = []

    for idx, student in enumerate(students, 1):
        stud_subject_marks = []
        row_values = [idx, student.roll_number, student.name]

        for sub in subjects:
            mark_obj = marks_map.get(student.id, {}).get(sub.id)
            is_absent = mark_obj.is_absent if mark_obj else False
            marks_val = mark_obj.marks_obtained if mark_obj else None

            stud_subject_marks.append({
                'subject_name': sub.name,
                'marks_obtained': marks_val,
                'out_of_marks': sub.out_of_marks,
                'is_absent': is_absent
            })

            m_display = "AB" if is_absent or marks_val is None else marks_val
            row_values.append(m_display)

        # Dynamic consolidated evaluation (FIX: Computer subject is NEVER omitted)
        con_res = calculate_student_consolidated(student.name, student.roll_number, stud_subject_marks, classroom.pass_percentage, grade_rules)
        class_student_results.append(con_res)

        row_values.extend([
            con_res['total_obtained'],
            con_res['total_out_of'],
            f"{con_res['overall_percentage']:.2f}%",
            con_res['overall_grade'],
            "PASS" if con_res['is_pass'] else ("ABSENT" if con_res['all_absent'] else "FAIL")
        ])

        ws_con.append(row_values)

        current_row = idx + 3
        for c_idx in range(1, len(row_values) + 1):
            cell = ws_con.cell(row=current_row, column=c_idx)
            cell.font = font_regular
            cell.border = thin_border
            if c_idx in [1, 2] or c_idx >= 4:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Consolidated Grade Distribution Count at bottom
    c_overview = calculate_class_overview(class_student_results, classroom.pass_percentage, grade_rules)
    c_analysis_row = len(students) + 6

    ws_con.cell(row=c_analysis_row, column=1, value="CLASS-WIDE OVERALL GRADE DISTRIBUTION").font = font_bold
    summary_headers = ['Grade A', 'Grade B', 'Grade C', 'Grade D', 'Grade E', 'Absent (AB)', 'Total Students', 'Class Avg %', 'Pass %']

    for c_idx, sh in enumerate(summary_headers, 1):
        cell = ws_con.cell(row=c_analysis_row + 1, column=c_idx)
        cell.value = sh
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    c_gc = c_overview['grade_counts']
    summary_values = [
        c_gc.get('A', 0), c_gc.get('B', 0), c_gc.get('C', 0), c_gc.get('D', 0), c_gc.get('E', 0), c_gc.get('AB', 0),
        c_overview['total_students'], f"{c_overview['class_avg_percentage']}%", f"{c_overview['overall_pass_percentage']}%"
    ]
    for c_idx, sv in enumerate(summary_values, 1):
        cell = ws_con.cell(row=c_analysis_row + 2, column=c_idx)
        cell.value = sv
        cell.font = font_bold
        cell.alignment = align_center
        cell.border = thin_border

    for col in ws_con.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_con.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # -------------------------------------------------------------
    # 3. OVERALL RESULT ANALYSIS SHEET
    # -------------------------------------------------------------
    ws_ov = wb.create_sheet(title="Overall Result Analysis")
    ws_ov.views.sheetView[0].showGridLines = True

    ws_ov.merge_cells('A1:I1')
    title_cell = ws_ov['A1']
    title_cell.value = f"{classroom.class_name} {classroom.division} ({classroom.academic_year.name}) - Overall Subject Grade Comparison Analysis"
    title_cell.font = font_title
    title_cell.fill = fill_header
    title_cell.alignment = align_center

    ov_headers = ['SUBJECT NAME', 'GRADE A', 'GRADE B', 'GRADE C', 'GRADE D', 'GRADE E', 'ABSENT (AB)', 'AVG %', 'PASS %']
    ws_ov.append([])
    ws_ov.append(ov_headers)

    for c_idx, oh in enumerate(ov_headers, 1):
        cell = ws_ov.cell(row=3, column=c_idx)
        cell.font = font_header
        cell.fill = fill_sub_header
        cell.alignment = align_center
        cell.border = thin_border

    for idx, sub in enumerate(subjects, 1):
        analysis = subject_analysis_results.get(sub.id, {})
        gc = analysis.get('grade_counts', {})
        ov_row = [
            sub.name,
            gc.get('A', 0),
            gc.get('B', 0),
            gc.get('C', 0),
            gc.get('D', 0),
            gc.get('E', 0),
            gc.get('AB', 0),
            f"{analysis.get('subject_avg_percentage', 0.0):.2f}%",
            f"{analysis.get('subject_pass_percent', 0.0):.2f}%",
        ]
        ws_ov.append(ov_row)

        current_row = idx + 3
        for c_idx in range(1, 10):
            cell = ws_ov.cell(row=current_row, column=c_idx)
            cell.font = font_regular
            cell.border = thin_border
            cell.alignment = align_left if c_idx == 1 else align_center

    for col in ws_ov.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_ov.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # Save to BytesIO stream
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
